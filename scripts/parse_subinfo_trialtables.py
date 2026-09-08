"""Extract per-session trial tables from the SubInfo2 experimenter workbook.

Each session sheet (s01_s02, s03_s04, ...) of SubInfo2*.xlsx carries an
"Experiment Trials" block: a header row (Trial #, rand #, Distance (m),
Package size, Hand-off pose, Rep1 status, Rep2 status, Notes) followed by 48
trial rows. This writes one trialtable_<DATE>.csv per sheet in the same
column order the earlier hand-made CSVs used, so
brock_functions.load_condition_table() reads them unchanged.

Usage:
    uv run python parse_subinfo_trialtables.py [XLSX] [OUT_DIR] [--sheets s05_s06,s07_s08]
"""
import argparse
import os
import re

import openpyxl
import pandas as pd

DEFAULT_XLSX = ('/Users/jeremy/Dropbox/Treadmill Brock 2025/imu data/'
                'SubInfo2(2).xlsx')

COLUMNS = ['Trial #', 'rand #', 'Distance (m)', 'Package size',
           'Hand-off pose', 'Rep1 status', 'Rep2 status', 'Notes']


def parse_sheet(ws):
    """Return (date_str, trials DataFrame) from one session sheet.

    Finds the 'Trial #' header cell, then reads every following row whose
    Trial # cell is an integer. The date comes from the 'Date' row of the
    subject-info block (column C), normalized to YYYYMMDD.
    """
    date = None
    header_pos = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == 'Date':
                raw = ws.cell(row=cell.row, column=cell.column + 1).value
                if raw is not None:
                    date = re.sub(r'\D', '', str(raw)[:10])
            if cell.value == 'Trial #':
                header_pos = (cell.row, cell.column)
    if header_pos is None:
        raise ValueError(f'sheet {ws.title!r}: no "Trial #" header found')

    r0, c0 = header_pos
    records = []
    for r in range(r0 + 1, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c0 + k).value for k in range(len(COLUMNS))]
        if not isinstance(vals[0], (int, float)):
            break
        records.append(vals)
    df = pd.DataFrame(records, columns=COLUMNS)
    df['Trial #'] = df['Trial #'].astype(int)
    return date, df


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('xlsx', nargs='?', default=DEFAULT_XLSX)
    ap.add_argument('out_dir', nargs='?', default=None,
                    help='default: directory of the xlsx')
    ap.add_argument('--sheets', default=None,
                    help='comma-separated sheet names (default: every '
                         'sheet with a filled trial block)')
    args = ap.parse_args()
    out_dir = args.out_dir or os.path.dirname(os.path.abspath(args.xlsx))
    wanted = args.sheets.split(',') if args.sheets else None

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    for name in wb.sheetnames:
        if wanted is not None and name not in wanted:
            continue
        if not re.match(r's\d+_s\d+', name):
            continue
        try:
            date, df = parse_sheet(wb[name])
        except ValueError as e:
            print(f'skip {name}: {e}')
            continue
        if not date or df.empty or df['Distance (m)'].isna().all():
            print(f'skip {name}: no date or empty trial block')
            continue
        out = os.path.join(out_dir, f'trialtable_{date}.csv')
        df.to_csv(out, index=False)
        n_notes = int((df['Rep1 status'].notna() | df['Rep2 status'].notna()).sum())
        print(f'{name} -> {out}  ({len(df)} trials, {n_notes} with rep notes)')


if __name__ == '__main__':
    main()
